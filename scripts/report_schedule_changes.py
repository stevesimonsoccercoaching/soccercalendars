from __future__ import annotations

import argparse
import json
import re
import subprocess
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_COLUMNS = {
    "Match #",
    "Date",
    "Time",
    "Home Team",
    "Away Team",
    "Location",
    "Division",
    "Status",
}

COMPARE_FIELDS = (
    ("date", "Date"),
    ("time", "Time"),
    ("opponent", "Opponent"),
    ("home_away", "Home/Away"),
    ("location", "Location"),
    ("division", "Division"),
    ("status", "Status"),
    ("result", "Results"),
)


def clean_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def run_git(args, *, binary=False, check=True):
    result = subprocess.run(
        ["git", *args],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
    )

    if check and result.returncode != 0:
        raise RuntimeError(
            f"git {' '.join(args)} failed:\n"
            f"{result.stderr.decode('utf-8', errors='replace')}"
        )

    if binary:
        return result.stdout

    return result.stdout.decode("utf-8", errors="replace")


def commit_exists(ref):
    if not ref or set(ref) == {"0"}:
        return False

    result = subprocess.run(
        ["git", "cat-file", "-e", f"{ref}^{{commit}}"],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        check=False,
    )
    return result.returncode == 0


def resolve_base(base, head):
    if commit_exists(base):
        return base

    parent = run_git(["rev-parse", f"{head}^"], check=False).strip()
    if commit_exists(parent):
        return parent

    return ""


def git_file_bytes(ref, path):
    result = subprocess.run(
        ["git", "show", f"{ref}:{path}"],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        check=False,
    )

    if result.returncode != 0:
        return None

    return result.stdout


def normalize_match_no(value):
    if value is None:
        return ""

    if isinstance(value, bool):
        return clean_text(value)

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    text = clean_text(value)
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]

    return text


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = clean_text(value)
    if not text:
        return None

    formats = (
        "%A, %B %d, %Y",
        "%B %d, %Y",
        "%m/%d/%Y",
        "%m/%d/%y",
        "%Y-%m-%d",
    )

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    return None


def date_values(value):
    parsed = parse_date(value)
    if parsed:
        return parsed.isoformat(), parsed.strftime("%a, %b %-d, %Y")

    text = clean_text(value)
    return text.casefold(), text or "(blank)"


def time_values(value):
    if isinstance(value, datetime):
        parsed = value.time()
    elif isinstance(value, time):
        parsed = value
    else:
        text = clean_text(value)
        if text.upper() in {
            "",
            "-",
            "TBD",
            "TBA",
            "TO BE DETERMINED",
            "TO BE ANNOUNCED",
        }:
            return "TBD", "TBD"

        match = re.match(r"^(\d{1,2}):(\d{2})", text)
        if not match:
            return text.casefold(), text

        hour = int(match.group(1))
        minute = int(match.group(2))
        parsed = time(hour=hour, minute=minute)

    canonical = parsed.strftime("%H:%M")
    display = parsed.strftime("%I:%M %p").lstrip("0")
    return canonical, display


def text_values(value, *, blank_display="(blank)", casefold=True):
    text = clean_text(value)
    canonical = text.casefold() if casefold else text
    return canonical, text or blank_display


def status_values(value):
    text = clean_text(value) or "Scheduled"
    return text.casefold(), text


def result_values(value):
    text = clean_text(value)

    if not text or text == "-":
        return "", "(none)"

    score = re.fullmatch(r"(\d+)\s*-\s*(\d+)", text)
    if score:
        canonical = f"{score.group(1)}-{score.group(2)}"
        display = f"{score.group(1)} - {score.group(2)}"
        return canonical, display

    return text.casefold(), text


def read_schedule(workbook_bytes, tracked_team_name=""):
    workbook = load_workbook(
        BytesIO(workbook_bytes),
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active
    header_row = next(sheet.iter_rows(values_only=True), None)

    if not header_row:
        workbook.close()
        raise RuntimeError("Spreadsheet has no header row")

    headers = {
        clean_text(value): index
        for index, value in enumerate(header_row)
        if value is not None
    }

    missing = REQUIRED_COLUMNS - set(headers)
    if missing:
        workbook.close()
        raise RuntimeError(
            f"Spreadsheet is missing required columns: {sorted(missing)}"
        )

    matches = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        match_no = normalize_match_no(row[headers["Match #"]])
        if not match_no:
            continue

        if match_no in matches:
            workbook.close()
            raise RuntimeError(f"Duplicate Match # {match_no}")

        home_team = clean_text(row[headers["Home Team"]])
        away_team = clean_text(row[headers["Away Team"]])

        date_canon, date_display = date_values(row[headers["Date"]])
        time_canon, time_display = time_values(row[headers["Time"]])
        location_canon, location_display = text_values(
            row[headers["Location"]]
        )
        division_canon, division_display = text_values(
            row[headers["Division"]]
        )
        status_canon, status_display = status_values(
            row[headers["Status"]]
        )

        if "Results" in headers:
            result_canon, result_display = result_values(
                row[headers["Results"]]
            )
        else:
            result_canon, result_display = "", "(none)"

        home_away = ""
        opponent = ""

        if tracked_team_name:
            if home_team == tracked_team_name:
                home_away = "HOME"
                opponent = away_team
            elif away_team == tracked_team_name:
                home_away = "AWAY"
                opponent = home_team

        if not opponent:
            opponent = f"{home_team} vs {away_team}"

        matches[match_no] = {
            "match_no": match_no,
            "home_team": home_team,
            "away_team": away_team,
            "date": date_canon,
            "date_display": date_display,
            "time": time_canon,
            "time_display": time_display,
            "opponent": clean_text(opponent).casefold(),
            "opponent_display": clean_text(opponent),
            "home_away": home_away,
            "home_away_display": home_away or "(not determined)",
            "location": location_canon,
            "location_display": location_display,
            "division": division_canon,
            "division_display": division_display,
            "status": status_canon,
            "status_display": status_display,
            "result": result_canon,
            "result_display": result_display,
        }

    workbook.close()
    return matches


def load_manifest_for_data_path(ref, data_path):
    path = Path(data_path)
    try:
        relative = path.relative_to("data")
    except ValueError:
        return {}, {}

    manifest_path = Path("config") / relative.parent / "manifest.json"
    raw = git_file_bytes(ref, manifest_path.as_posix())
    if not raw:
        return {}, {}

    manifest = json.loads(raw.decode("utf-8"))
    slug = path.stem

    team = next(
        (item for item in manifest.get("teams", []) if item.get("slug") == slug),
        {},
    )
    return manifest.get("competition", {}), team


def file_context(head, base, path):
    competition, team = load_manifest_for_data_path(head, path)
    if not team:
        competition, team = load_manifest_for_data_path(base, path)

    organizer = clean_text(competition.get("organizer"))
    season = clean_text(competition.get("season"))
    year = clean_text(competition.get("year"))

    competition_label = " ".join(
        part for part in (organizer, season, year) if part
    ) or "Schedule"

    team_name = clean_text(team.get("name"))
    short_name = clean_text(team.get("short_name")) or team_name or Path(path).stem

    return {
        "competition_label": competition_label,
        "team_name": team_name,
        "short_name": short_name,
    }


def match_sort_key(match_no):
    if match_no.isdigit():
        return (0, int(match_no))
    return (1, match_no.casefold())


def render_match_snapshot(match):
    lines = [
        f"- Date: **{match['date_display']}**",
        f"- Time: **{match['time_display']}**",
        f"- HOME: {match['home_team'] or '(blank)'}",
        f"- AWAY: {match['away_team'] or '(blank)'}",
        f"- Location: {match['location_display']}",
        f"- Division: {match['division_display']}",
        f"- Status: {match['status_display']}",
    ]

    if match["result"]:
        lines.append(f"- Results: {match['result_display']}")

    return lines


def render_modified_match(old_match, new_match):
    changes = []

    for key, label in COMPARE_FIELDS:
        if old_match[key] == new_match[key]:
            continue

        old_display = old_match[f"{key}_display"]
        new_display = new_match[f"{key}_display"]
        changes.append(
            f"- {label}: **{old_display}** → **{new_display}**"
        )

    home_changed = old_match["home_team"].casefold() != new_match["home_team"].casefold()
    away_changed = old_match["away_team"].casefold() != new_match["away_team"].casefold()

    if not old_match["home_away"] and not new_match["home_away"]:
        if home_changed:
            changes.append(
                f"- HOME team: **{old_match['home_team']}** → "
                f"**{new_match['home_team']}**"
            )
        if away_changed:
            changes.append(
                f"- AWAY team: **{old_match['away_team']}** → "
                f"**{new_match['away_team']}**"
            )

    return changes


def parse_changed_files(base, head):
    output = run_git(
        ["diff", "--name-status", "-M", base, head, "--", "data"]
    )

    changes = []

    for line in output.splitlines():
        if not line.strip():
            continue

        parts = line.split("\t")
        status = parts[0]

        if status.startswith("R") and len(parts) >= 3:
            old_path, new_path = parts[1], parts[2]
            if old_path.endswith(".xlsx") or new_path.endswith(".xlsx"):
                changes.append(("R", old_path, new_path))
            continue

        if len(parts) < 2:
            continue

        path = parts[1]
        if not path.endswith(".xlsx"):
            continue

        changes.append((status[0], path, path))

    return changes


def build_report(base, head):
    short_base = base[:7] if base else "none"
    short_head = head[:7]

    lines = [
        "# Soccer schedule change report",
        "",
        f"Comparison: `{short_base}` → `{short_head}`",
        "",
        (
            "Rows are compared by GotSport **Match #**. A blank Status is treated "
            "as **Scheduled**, and a blank/`-` Results value is treated as no result."
        ),
        "",
    ]

    if not base:
        lines.extend(
            [
                "No usable previous commit was available, so there is nothing to compare.",
                "",
            ]
        )
        return "\n".join(lines)

    changed_files = parse_changed_files(base, head)

    if not changed_files:
        lines.extend(
            [
                "## No XLSX schedule files changed",
                "",
                "The calendar build can continue normally.",
                "",
            ]
        )
        return "\n".join(lines)

    total_added = 0
    total_removed = 0
    total_modified = 0
    files_with_material_changes = 0

    for status, old_path, new_path in changed_files:
        display_path = new_path if status != "D" else old_path
        context = file_context(head, base, display_path)
        tracked_team = context["team_name"]

        old_bytes = git_file_bytes(base, old_path) if status != "A" else None
        new_bytes = git_file_bytes(head, new_path) if status != "D" else None

        old_matches = read_schedule(old_bytes, tracked_team) if old_bytes else {}
        new_matches = read_schedule(new_bytes, tracked_team) if new_bytes else {}

        lines.extend(
            [
                f"## {context['short_name']}",
                "",
                f"Competition: **{context['competition_label']}**  ",
                f"File: `{display_path}`",
                "",
            ]
        )

        if status == "A":
            total_added += len(new_matches)
            files_with_material_changes += 1
            lines.append(
                f"**New schedule file added with {len(new_matches)} match(es).**"
            )
            lines.append("")
            continue

        if status == "D":
            total_removed += len(old_matches)
            files_with_material_changes += 1
            lines.append(
                f"**Schedule file removed; it previously contained "
                f"{len(old_matches)} match(es).**"
            )
            lines.append("")
            continue

        if status == "R" and old_path != new_path:
            lines.extend(
                [
                    f"File renamed from `{old_path}` to `{new_path}`.",
                    "",
                ]
            )

        old_ids = set(old_matches)
        new_ids = set(new_matches)
        added_ids = sorted(new_ids - old_ids, key=match_sort_key)
        removed_ids = sorted(old_ids - new_ids, key=match_sort_key)
        common_ids = sorted(old_ids & new_ids, key=match_sort_key)

        modified = []
        for match_no in common_ids:
            change_lines = render_modified_match(
                old_matches[match_no],
                new_matches[match_no],
            )
            if change_lines:
                modified.append((match_no, change_lines))

        if not added_ids and not removed_ids and not modified:
            lines.extend(
                [
                    "The XLSX file changed, but no material schedule fields changed.",
                    "",
                ]
            )
            continue

        files_with_material_changes += 1
        total_added += len(added_ids)
        total_removed += len(removed_ids)
        total_modified += len(modified)

        for match_no in added_ids:
            lines.extend(
                [
                    f"### Match {match_no} — ADDED",
                    "",
                    *render_match_snapshot(new_matches[match_no]),
                    "",
                ]
            )

        for match_no in removed_ids:
            lines.extend(
                [
                    f"### Match {match_no} — REMOVED",
                    "",
                    *render_match_snapshot(old_matches[match_no]),
                    "",
                ]
            )

        for match_no, change_lines in modified:
            lines.extend(
                [
                    f"### Match {match_no} — CHANGED",
                    "",
                    *change_lines,
                    "",
                ]
            )

    lines.extend(
        [
            "---",
            "",
            "## Summary",
            "",
            f"- XLSX files changed in Git: **{len(changed_files)}**",
            f"- Files with material schedule changes: **{files_with_material_changes}**",
            f"- Matches added: **{total_added}**",
            f"- Matches removed: **{total_removed}**",
            f"- Existing matches changed: **{total_modified}**",
            "",
        ]
    )

    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(
        description="Compare changed soccer XLSX schedules between two Git commits."
    )
    parser.add_argument("--base", default="")
    parser.add_argument("--head", default="HEAD")
    parser.add_argument("--output", default="schedule-change-report.md")
    args = parser.parse_args()

    head = run_git(["rev-parse", args.head]).strip()
    base = resolve_base(args.base, head)

    report = build_report(base, head)
    output_file = Path(args.output)
    output_file.write_text(report + "\n", encoding="utf-8")

    print(report)
    print(f"\nSaved report to {output_file}")


if __name__ == "__main__":
    main()

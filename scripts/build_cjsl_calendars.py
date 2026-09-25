from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
import json
import re
import sys

from openpyxl import load_workbook


REQUIRED_COLUMNS = {
    "Match #",
    "Date",
    "Time",
    "Home Team",
    "Away Team",
    "Location",
    "Division",
    "Status",
}


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def ics_escape(value):
    value = clean_text(value)
    return (
        value
        .replace("\\", "\\\\")
        .replace(";", "\\;")
        .replace(",", "\\,")
        .replace("\r\n", "\\n")
        .replace("\n", "\\n")
    )


def fold_line(line, limit=75):
    if len(line.encode("utf-8")) <= limit:
        return [line]

    pieces = []
    remaining = line
    first = True

    while remaining:
        available_bytes = limit if first else limit - 1
        piece = ""
        used_bytes = 0

        for character in remaining:
            character_bytes = len(character.encode("utf-8"))
            if used_bytes + character_bytes > available_bytes:
                break
            piece += character
            used_bytes += character_bytes

        if not piece:
            piece = remaining[0]

        pieces.append(piece if first else " " + piece)
        remaining = remaining[len(piece):]
        first = False

    return pieces


def parse_match_date(date_value):
    if isinstance(date_value, datetime):
        return date_value.date()

    if isinstance(date_value, date):
        return date_value

    date_text = clean_text(date_value)

    if not date_text:
        raise ValueError("Match date is blank")

    date_text = re.sub(r"\s+", " ", date_text)

    return datetime.strptime(
        date_text,
        "%A, %B %d, %Y"
    ).date()


def parse_match_start(date_value, time_value):
    match_date = parse_match_date(date_value)

    if isinstance(time_value, datetime):
        match_time = time_value.time()
    elif isinstance(time_value, time):
        match_time = time_value
    else:
        time_text = clean_text(time_value)

        if not time_text:
            raise ValueError("Match time is blank")

        clock_text = time_text.split()[0]
        match_time = datetime.strptime(
            clock_text,
            "%H:%M"
        ).time()

    return datetime.combine(match_date, match_time)


def time_is_tbd(time_value):
    if isinstance(time_value, (datetime, time)):
        return False

    time_text = clean_text(time_value).upper()

    return time_text in {
        "",
        "-",
        "TBD",
        "TBA",
        "TO BE DETERMINED",
        "TO BE ANNOUNCED",
    }


def read_matches(input_file, tracked_team_name, duration_minutes):
    workbook = load_workbook(
        input_file,
        read_only=True,
        data_only=True
    )

    sheet = workbook.active
    header_row = next(sheet.iter_rows(values_only=True))

    headers = {
        clean_text(value): index
        for index, value in enumerate(header_row)
        if value is not None
    }

    missing = REQUIRED_COLUMNS - set(headers)

    if missing:
        raise RuntimeError(
            f"{input_file} is missing required columns: {sorted(missing)}"
        )

    matches = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        match_no = row[headers["Match #"]]

        if match_no is None:
            continue

        home_team = clean_text(row[headers["Home Team"]])
        away_team = clean_text(row[headers["Away Team"]])

        if home_team != tracked_team_name and away_team != tracked_team_name:
            raise RuntimeError(
                f"Match {match_no} does not contain configured team "
                f"{tracked_team_name!r}"
            )

        if "Results" in headers:
            result = clean_text(row[headers["Results"]])
        else:
            result = ""

        raw_status = clean_text(row[headers["Status"]])
        display_status = raw_status if raw_status else "Scheduled"

        date_value = row[headers["Date"]]
        time_value = row[headers["Time"]]

        try:
            match_date = parse_match_date(date_value)
        except ValueError as exc:
            raise RuntimeError(
                f"{input_file}: Match {match_no} has no usable date"
            ) from exc

        all_day = time_is_tbd(time_value)

        if all_day:
            start = datetime.combine(match_date, time.min)
            end = start + timedelta(days=1)
        else:
            try:
                start = parse_match_start(date_value, time_value)
            except (ValueError, IndexError) as exc:
                raise RuntimeError(
                    f"{input_file}: Match {match_no} has an invalid "
                    f"time value: {time_value!r}"
                ) from exc

            end = start + timedelta(minutes=duration_minutes)

        matches.append(
            {
                "match_no": int(match_no),
                "start": start,
                "end": end,
                "all_day": all_day,
                "home_team": home_team,
                "away_team": away_team,
                "result": result,
                "location": clean_text(row[headers["Location"]]),
                "division": clean_text(row[headers["Division"]]),
                "status": display_status,
            }
        )

    workbook.close()

    matches.sort(
        key=lambda match: (
            match["start"],
            match["match_no"]
        )
    )

    return matches


def build_calendar(
    competition,
    defaults,
    display_names,
    team,
    matches,
    output_file
):
    organizer = competition["organizer"]
    event_id = str(competition["event_id"])
    calendar_timezone = competition["timezone"]

    team_name = team["name"]
    short_name = team.get(
        "short_name",
        display_names.get(team_name, team_name)
    )

    source_url = team["gotsport"]["schedule_url"]
    reminders = defaults.get("reminders", [])

    generated_at = datetime.now(timezone.utc).strftime(
        "%Y%m%dT%H%M%SZ"
    )

    calendar_name = f"{organizer} - {team_name}"

    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//Soccer Calendars//CJSL GotSport Converter//EN",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        f"NAME:{ics_escape(calendar_name)}",
        f"X-WR-CALNAME:{ics_escape(calendar_name)}",
        f"X-WR-TIMEZONE:{calendar_timezone}",
        "BEGIN:VTIMEZONE",
        f"TZID:{calendar_timezone}",
        f"X-LIC-LOCATION:{calendar_timezone}",
        "BEGIN:DAYLIGHT",
        "TZOFFSETFROM:-0500",
        "TZOFFSETTO:-0400",
        "TZNAME:EDT",
        "DTSTART:19700308T020000",
        "RRULE:FREQ=YEARLY;BYMONTH=3;BYDAY=2SU",
        "END:DAYLIGHT",
        "BEGIN:STANDARD",
        "TZOFFSETFROM:-0400",
        "TZOFFSETTO:-0500",
        "TZNAME:EST",
        "DTSTART:19701101T020000",
        "RRULE:FREQ=YEARLY;BYMONTH=11;BYDAY=1SU",
        "END:STANDARD",
        "END:VTIMEZONE",
    ]

    for match in matches:
        home_team = match["home_team"]
        away_team = match["away_team"]

        if home_team == team_name:
            home_away = "HOME"
            opponent_full = away_team
            separator = "vs"
        else:
            home_away = "AWAY"
            opponent_full = home_team
            separator = "@"

        opponent_short = display_names.get(
            opponent_full,
            opponent_full
        )

        summary = (
            f"{organizer}: {short_name} "
            f"{separator} {opponent_short}"
        )

        description_parts = [
            f"Status: {ics_escape(match['status'])}",
            f"Match No: {match['match_no']} ({home_away})",
            f"Division: {ics_escape(match['division'])}",
            "",
            f"HOME: {ics_escape(home_team)}",
            f"AWAY: {ics_escape(away_team)}",
        ]

        result = clean_text(match.get("result", ""))

        if result and result != "-":
            description_parts.extend(
                [
                    "",
                    f"Results: {ics_escape(result)}",
                ]
            )

        description_parts.extend(
            [
                "",
                f"Source: {source_url}",
            ]
        )

        description = "\\n".join(description_parts)

        uid = (
            f"{organizer.lower()}-{event_id}-"
            f"match-{match['match_no']}@system.gotsport.com"
        )

        if match["all_day"]:
            dtstart_line = (
                "DTSTART;VALUE=DATE:"
                f"{match['start'].strftime('%Y%m%d')}"
            )
            dtend_line = (
                "DTEND;VALUE=DATE:"
                f"{match['end'].strftime('%Y%m%d')}"
            )
        else:
            dtstart_line = (
                f"DTSTART;TZID={calendar_timezone}:"
                f"{match['start'].strftime('%Y%m%dT%H%M%S')}"
            )
            dtend_line = (
                f"DTEND;TZID={calendar_timezone}:"
                f"{match['end'].strftime('%Y%m%dT%H%M%S')}"
            )

        lines.extend(
            [
                "BEGIN:VEVENT",
                f"UID:{uid}",
                f"DTSTAMP:{generated_at}",
                dtstart_line,
                dtend_line,
                f"SUMMARY:{ics_escape(summary)}",
                f"LOCATION:{ics_escape(match['location'])}",
                f"DESCRIPTION:{description}",
                f"URL:{source_url}",
                f"CATEGORIES:{organizer},Soccer",
                "TRANSP:OPAQUE",
            ]
        )

        status_lower = match["status"].lower()

        if "cancel" in status_lower:
            lines.append("STATUS:CANCELLED")
        else:
            lines.append("STATUS:CONFIRMED")

        lines.append(
            f"X-GOTSPORT-STATUS:{ics_escape(match['status'])}"
        )

        if not match["all_day"]:
            for reminder in reminders:
                lines.extend(
                    [
                        "BEGIN:VALARM",
                        f"TRIGGER:{reminder['trigger']}",
                        "ACTION:DISPLAY",
                        f"DESCRIPTION:{ics_escape(reminder['description'])}",
                        "END:VALARM",
                    ]
                )

        lines.append("END:VEVENT")

    lines.append("END:VCALENDAR")

    folded_lines = []

    for line in lines:
        folded_lines.extend(fold_line(line))

    output_file.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    output_file.write_text(
        "\r\n".join(folded_lines) + "\r\n",
        encoding="utf-8",
        newline=""
    )

    print(
        f"Created {output_file} with {len(matches)} matches."
    )


def main():
    if len(sys.argv) != 2:
        raise SystemExit(
            "Usage:\n"
            "python scripts/build_cjsl_calendars.py "
            "config/YEAR/cjsl/SEASON/manifest.json"
        )

    manifest_file = Path(sys.argv[1])

    if not manifest_file.exists():
        raise SystemExit(
            f"Manifest not found: {manifest_file}"
        )

    manifest = json.loads(
        manifest_file.read_text(encoding="utf-8")
    )

    competition = manifest["competition"]
    defaults = manifest.get("calendar_defaults", {})
    display_names = manifest.get("display_names", {})
    teams = manifest.get("teams", [])

    if not teams:
        raise SystemExit(
            "No teams are configured in the manifest."
        )

    parts = manifest_file.parts

    try:
        config_index = parts.index("config")
    except ValueError:
        raise SystemExit(
            "Manifest must be stored beneath config/."
        )

    relative_folder = Path(
        *parts[config_index + 1:-1]
    )

    data_folder = Path("data") / relative_folder
    docs_folder = Path("docs") / relative_folder

    default_duration = int(
        competition.get(
            "default_duration_minutes",
            60
        )
    )

    print(
        f"Building {len(teams)} calendar(s) for "
        f"{competition['name']}..."
    )

    for team in teams:
        team_name = team["name"]
        slug = team["slug"]

        input_file = data_folder / f"{slug}.xlsx"
        output_file = docs_folder / f"{slug}.ics"

        if not input_file.exists():
            raise FileNotFoundError(
                f"Missing XLSX for {team_name}:\n"
                f"{input_file}"
            )

        duration_minutes = int(
            team.get(
                "duration_minutes",
                default_duration
            )
        )

        matches = read_matches(
            input_file=input_file,
            tracked_team_name=team_name,
            duration_minutes=duration_minutes,
        )

        build_calendar(
            competition=competition,
            defaults=defaults,
            display_names=display_names,
            team=team,
            matches=matches,
            output_file=output_file,
        )

    print(
        "All configured CJSL calendars built successfully."
    )


if __name__ == "__main__":
    main()

from collections import Counter, defaultdict
from datetime import date, datetime
from html import escape
from pathlib import Path
import json
import re

from openpyxl import load_workbook


CONFIG_ROOT = Path("config")
DATA_ROOT = Path("data")
DOCS_ROOT = Path("docs")
OUTPUT_FILE = DOCS_ROOT / "index.html"

SEASON_FALLBACK_MONTH = {
    "winter": 1,
    "spring": 4,
    "summer": 7,
    "fall": 10,
    "autumn": 10,
}


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_date_value(value):
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = clean_text(value)
    if not text:
        return None

    text = re.sub(r"\s+", " ", text)

    formats = (
        "%A, %B %d, %Y",
        "%B %d, %Y",
        "%m/%d/%Y",
        "%m/%d/%y",
        "%Y-%m-%d",
    )

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue

    return None


def natural_key(value):
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", clean_text(value))
    ]


def season_sort_date(competition, schedule_dates):
    if schedule_dates:
        return max(schedule_dates)

    year = int(competition.get("year", 0) or 0)
    season = clean_text(competition.get("season", "")).lower()
    month = SEASON_FALLBACK_MONTH.get(season, 1)

    if year:
        return date(year, month, 1)

    return date.min


def display_division(raw_division, competition):
    division = clean_text(raw_division) or "Division not listed"

    if re.match(r"^(male|female|boys?|girls?|coed|mixed)\b", division, re.I):
        return division

    gender = clean_text(competition.get("gender", "Male"))

    if not gender:
        return division

    return f"{gender} {division}"


def read_team_schedule(xlsx_file):
    workbook = load_workbook(
        xlsx_file,
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active
    header_row = next(sheet.iter_rows(values_only=True), None)

    if not header_row:
        workbook.close()
        return {
            "division": "Division not listed",
            "dates": [],
        }

    headers = {
        clean_text(value): index
        for index, value in enumerate(header_row)
        if value is not None
    }

    division_values = []
    schedule_dates = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if "Division" in headers:
            division = clean_text(row[headers["Division"]])
            if division:
                division_values.append(division)

        if "Date" in headers:
            match_date = parse_date_value(row[headers["Date"]])
            if match_date:
                schedule_dates.append(match_date)

    workbook.close()

    if division_values:
        counts = Counter(division_values)
        division = counts.most_common(1)[0][0]

        if len(counts) > 1:
            print(
                f"WARNING: {xlsx_file} contains multiple divisions: "
                f"{sorted(counts)}. Using {division!r} for the index."
            )
    else:
        division = "Division not listed"

    return {
        "division": division,
        "dates": schedule_dates,
    }


def load_competitions():
    manifests = sorted(CONFIG_ROOT.glob("**/manifest.json"))
    competitions = []

    for manifest_file in manifests:
        manifest = json.loads(
            manifest_file.read_text(encoding="utf-8")
        )

        competition = manifest.get("competition", {})
        teams = manifest.get("teams", [])

        if not teams:
            print(f"Skipping {manifest_file}: no teams configured.")
            continue

        relative_folder = manifest_file.parent.relative_to(CONFIG_ROOT)
        data_folder = DATA_ROOT / relative_folder
        docs_folder = DOCS_ROOT / relative_folder

        team_records = []
        competition_dates = []

        for team in teams:
            slug = clean_text(team.get("slug"))
            official_name = clean_text(team.get("name"))
            short_name = clean_text(team.get("short_name")) or official_name

            if not slug:
                print(
                    f"WARNING: Skipping team without slug in {manifest_file}: "
                    f"{official_name or '<unnamed>'}"
                )
                continue

            xlsx_file = data_folder / f"{slug}.xlsx"
            ics_file = docs_folder / f"{slug}.ics"

            if not xlsx_file.exists():
                print(
                    f"WARNING: Skipping {official_name}: missing XLSX "
                    f"{xlsx_file}"
                )
                continue

            if not ics_file.exists():
                print(
                    f"WARNING: Skipping {official_name}: missing generated ICS "
                    f"{ics_file}"
                )
                continue

            schedule = read_team_schedule(xlsx_file)
            competition_dates.extend(schedule["dates"])

            href = ics_file.relative_to(DOCS_ROOT).as_posix()

            team_records.append(
                {
                    "name": official_name,
                    "short_name": short_name,
                    "slug": slug,
                    "division": display_division(
                        schedule["division"],
                        competition,
                    ),
                    "href": href,
                }
            )

        if not team_records:
            print(
                f"Skipping {manifest_file}: no teams have both XLSX and ICS files."
            )
            continue

        organizer = clean_text(competition.get("organizer")) or "Other"
        season = clean_text(competition.get("season"))
        year = competition.get("year", "")

        if season and year:
            season_label = f"{season} {year}"
        elif clean_text(competition.get("name")):
            season_label = clean_text(competition.get("name"))
        else:
            season_label = "Competition"

        start_date = min(competition_dates) if competition_dates else None
        end_date = max(competition_dates) if competition_dates else None

        competitions.append(
            {
                "organizer": organizer,
                "season_label": season_label,
                "sort_date": season_sort_date(
                    competition,
                    competition_dates,
                ),
                "start_date": start_date,
                "end_date": end_date,
                "teams": team_records,
            }
        )

    return competitions


def render_team_list(teams):
    short_name_counts = Counter(
        team["short_name"].casefold()
        for team in teams
    )

    items = []

    for team in sorted(
        teams,
        key=lambda item: (
            natural_key(item["short_name"]),
            natural_key(item["name"]),
        ),
    ):
        if short_name_counts[team["short_name"].casefold()] > 1:
            display_name = team["name"]
        else:
            display_name = team["short_name"]

        items.append(
            "          <li>"
            f'<a class="calendar-link" href="{escape(team["href"], quote=True)}" '
            f'download title="{escape(team["name"], quote=True)}">'
            f"{escape(display_name)}"
            "</a></li>"
        )

    return "\n".join(items)


def render_competition(comp):
    divisions = defaultdict(list)

    for team in comp["teams"]:
        divisions[team["division"]].append(team)

    end_date = (
        comp["end_date"].isoformat()
        if comp["end_date"]
        else ""
    )

    start_date = (
        comp["start_date"].isoformat()
        if comp["start_date"]
        else ""
    )

    parts = [
        (
            f'    <section class="competition" '
            f'data-organizer="{escape(comp["organizer"], quote=True)}" '
            f'data-start-date="{start_date}" '
            f'data-end-date="{end_date}">'
        ),
        f"      <h3>{escape(comp['season_label'])}</h3>",
    ]

    for division in sorted(divisions, key=natural_key):
        parts.extend(
            [
                '      <div class="division">',
                f"        <h4>{escape(division)}</h4>",
                '        <ul class="team-list">',
                render_team_list(divisions[division]),
                "        </ul>",
                "      </div>",
            ]
        )

    parts.append("    </section>")
    return "\n".join(parts)


def build_index(competitions):
    organizers = defaultdict(list)

    for competition in competitions:
        organizers[competition["organizer"]].append(competition)

    organizer_sections = []

    for organizer in sorted(organizers, key=natural_key):
        comps = sorted(
            organizers[organizer],
            key=lambda item: item["sort_date"],
            reverse=True,
        )

        rendered_competitions = "\n".join(
            render_competition(comp)
            for comp in comps
        )

        organizer_sections.append(
            "\n".join(
                [
                    '<section class="league">',
                    f"  <h2>{escape(organizer)}</h2>",
                    rendered_competitions,
                    "</section>",
                ]
            )
        )

    body = "\n\n".join(organizer_sections)

    return f'''<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Soccer Calendars</title>
  <style>
    :root {{
      color-scheme: light dark;
      font-family: system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
    }}

    body {{
      margin: 0;
      line-height: 1.45;
      background: Canvas;
      color: CanvasText;
    }}

    main {{
      width: min(860px, calc(100% - 32px));
      margin: 0 auto;
      padding: 32px 0 64px;
    }}

    h1 {{
      margin: 0 0 8px;
      font-size: clamp(2rem, 8vw, 3rem);
      line-height: 1.05;
    }}

    .intro {{
      margin: 0 0 24px;
      max-width: 720px;
      color: color-mix(in srgb, CanvasText 75%, transparent);
    }}

    .history-control {{
      display: none;
      align-items: center;
      gap: 10px;
      margin: 22px 0 30px;
      font-weight: 650;
      cursor: pointer;
      user-select: none;
    }}

    .history-control.visible {{
      display: flex;
    }}

    .history-control input {{
      width: 20px;
      height: 20px;
      margin: 0;
    }}

    .league {{
      margin-top: 34px;
    }}

    h2 {{
      margin: 0;
      padding-bottom: 8px;
      border-bottom: 3px solid currentColor;
      font-size: 1.75rem;
    }}

    .competition {{
      margin-top: 28px;
    }}

    .competition.historical {{
      display: none;
    }}

    body.show-history .competition.historical {{
      display: block;
    }}

    h3 {{
      margin: 0 0 22px;
      font-size: 1.4rem;
    }}

    .division {{
      margin: 0 0 28px;
    }}

    h4 {{
      margin: 0 0 8px;
      padding-bottom: 6px;
      border-bottom: 1px solid color-mix(in srgb, CanvasText 35%, transparent);
      font-size: 1rem;
      font-weight: 750;
    }}

    .team-list {{
      list-style: none;
      margin: 0;
      padding: 0;
    }}

    .team-list li {{
      border-bottom: 1px solid color-mix(in srgb, CanvasText 12%, transparent);
    }}

    .calendar-link {{
      display: block;
      padding: 12px 4px;
      color: LinkText;
      font-weight: 650;
      text-decoration-thickness: 1px;
      text-underline-offset: 3px;
    }}

    .calendar-link:hover,
    .calendar-link:focus-visible {{
      text-decoration-thickness: 2px;
    }}

    .empty-message {{
      margin-top: 32px;
    }}

    @media (max-width: 520px) {{
      main {{
        width: min(100% - 24px, 860px);
        padding-top: 24px;
      }}

      .calendar-link {{
        padding: 14px 2px;
      }}
    }}
  </style>
</head>
<body>
<main>
  <h1>Soccer Calendars</h1>
  <p class="intro">
    Tap a team to download or open its ICS calendar feed. To subscribe in Google Calendar,
    long-press or copy the team link and add it under <strong>Other calendars → From URL</strong>.
  </p>

  <label class="history-control" id="history-control">
    <input type="checkbox" id="show-history">
    <span>Show previous seasons</span>
  </label>

{body if body else '  <p class="empty-message">No calendar feeds are available yet.</p>'}
</main>

<script>
(function () {{
  const sections = Array.from(document.querySelectorAll('.competition'));
  const toggle = document.getElementById('show-history');
  const control = document.getElementById('history-control');

  if (!sections.length) return;

  const today = new Date();
  today.setHours(0, 0, 0, 0);

  const byOrganizer = new Map();

  for (const section of sections) {{
    const organizer = section.dataset.organizer || '';
    if (!byOrganizer.has(organizer)) byOrganizer.set(organizer, []);
    byOrganizer.get(organizer).push(section);
  }}

  let historicalCount = 0;

  for (const group of byOrganizer.values()) {{
    const activeOrFuture = group.filter((section) => {{
      const value = section.dataset.endDate;
      if (!value) return true;
      const endDate = new Date(value + 'T00:00:00');
      return endDate >= today;
    }});

    if (activeOrFuture.length) {{
      for (const section of group) {{
        const value = section.dataset.endDate;
        if (!value) continue;
        const endDate = new Date(value + 'T00:00:00');
        if (endDate < today) {{
          section.classList.add('historical');
          historicalCount += 1;
        }}
      }}
    }} else {{
      // If every season for a league is already over, leave the newest one visible
      // and treat only the older ones as historical.
      group.slice(1).forEach((section) => {{
        section.classList.add('historical');
        historicalCount += 1;
      }});
    }}
  }}

  if (historicalCount > 0) {{
    control.classList.add('visible');
  }}

  toggle.addEventListener('change', function () {{
    document.body.classList.toggle('show-history', toggle.checked);
  }});
}})();
</script>
</body>
</html>
'''


def main():
    competitions = load_competitions()

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_FILE.write_text(
        build_index(competitions),
        encoding="utf-8",
        newline="\n",
    )

    team_count = sum(
        len(competition["teams"])
        for competition in competitions
    )

    print(
        f"Created {OUTPUT_FILE} from {len(competitions)} competition(s) "
        f"and {team_count} team calendar(s)."
    )


if __name__ == "__main__":
    main()
